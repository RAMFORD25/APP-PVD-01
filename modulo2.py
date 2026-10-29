"""Módulo 2 - VALIDACIÓN DE TRAMOS PARA SUSCRIPCIÓN DE CONVENIOS.

Flujo:  Gobierno remite lista -> Especialista zonal completa los tramos -> la envía a Sede Central ->
        Sede Central aprueba u observa -> con la lista aprobada se descarga el Anexo I (I-A rutinario / I-B periódico)
        y su resumen, usando la plantilla original (plantilla_anexo_i.xlsx).
"""
import io
import json
import random
import re
from copy import copy
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import gspread
import pandas as pd
import streamlit as st

TITULO = "VALIDACIÓN DE TRAMOS PARA SUSCRIPCIÓN DE CONVENIOS"
PLANTILLA = Path(__file__).parent / "plantilla_anexo_i.xlsx"
HOJA_LOTES, HOJA_TRAMOS = "Lotes", "Tramos"

GASTOS_PCT = Decimal("0.10")  # gastos operativos / supervisión = 10 % de la ejecución
TIPOS_GOBIERNO = ["Local", "Regional"]
MANTENIMIENTOS = ["Rutinario", "Periódico"]
REGIONES_NATURALES = ["COSTA", "SIERRA", "SELVA"]
ESTADOS_CONSERVACION = ["BUENO", "REGULAR", "MALO"]
SUPERFICIES_NO_PAV = ["AFIRMADO", "TROCHA", "SUELO NATURAL"]
SUPERFICIES_PAV = ["ASFALTADO", "PAVIMENTO RÍGIDO", "TRATAMIENTO SUPERFICIAL"]
UBIGEO_DEPTO = {"TACNA": "23", "MOQUEGUA": "18", "HUANCAVELICA": "09"}
ANIO_DEFECTO = {2026: "Año de la recuperación y consolidación de la economía peruana"}

ESTADO_ELAB, ESTADO_ENVIADO, ESTADO_OBS, ESTADO_APROB, ESTADO_ANUL = (
    "En elaboración", "Enviado", "Observado", "Aprobado", "Anulado")

COLS_TRAMO = ["codigo_ruta", "distrito", "ubigeo", "denom_sinac", "denom_interv", "region_natural", "provincia",
              "zona", "prog_ini", "x_ini", "y_ini", "prog_fin", "x_fin", "y_fin", "long_parcial", "long_total",
              "estado_cons", "superficie", "ancho", "costo", "plazo", "observaciones"]
NUM_COLS = ["zona", "x_ini", "y_ini", "x_fin", "y_fin", "long_parcial", "long_total", "ancho", "costo", "plazo"]
LOTES_COLS = ["lote_id", "tipo_gobierno", "pliego", "mantenimiento", "anio", "departamento", "provincia", "ht", "ley",
              "estado", "creado_por", "fecha_creacion", "fecha_envio", "revisado_por", "fecha_revision",
              "obs_central", "actualizado"]
TRAMOS_COLS = ["lote_id", "activo", "nro"] + COLS_TRAMO


# ---------------------------------------------------------------- utilidades
def ahora():
    return datetime.now(timezone.utc) - timedelta(hours=5)  # hora de Lima


def ts():
    return ahora().strftime("%Y-%m-%d %H:%M:%S")


def vacio(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(v, str) and not v.strip()


def num(v):
    if vacio(v):
        return None
    try:
        return Decimal(str(v).strip().replace(",", "."))
    except InvalidOperation:
        return None


def redondear(x):
    return int(Decimal(x).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def parse_prog(v):
    """Metros (int) o None. Acepta '5+498', '05+498', 'KM 5+498', 5498 (metros) o 5.498 (km)."""
    if vacio(v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
        return int(v)
    s = str(v).upper().replace("KM", "").replace(" ", "").replace(",", ".")
    m = re.fullmatch(r"(\d+)\+(\d{3})", s)
    if m:
        return int(m.group(1)) * 1000 + int(m.group(2))
    if re.fullmatch(r"\d+", s):
        return int(s)
    if re.fullmatch(r"\d+\.\d+", s):
        return redondear(Decimal(s) * 1000)
    return None


def fmt_prog(m):
    return f"{m // 1000:02d}+{m % 1000:03d}"


def es_pavimentado(superficie):
    return str(superficie).strip().upper() in SUPERFICIES_PAV


# ---------------------------------------------------------------- cálculo
def calcular_tramo(r, mantenimiento):
    """Aplica las fórmulas de Ejecución:
       Rutinario: ROUND(Longitud total x Costo Km/Mes x Plazo/30, 0)   Periódico: ROUND(Longitud x Costo Km, 0)."""
    ini, fin = parse_prog(r.get("prog_ini")), parse_prog(r.get("prog_fin"))
    por_prog = Decimal(fin - ini) / 1000 if ini is not None and fin is not None and fin > ini else None
    parcial = num(r.get("long_parcial"))
    parcial = parcial if parcial is not None else por_prog
    total = num(r.get("long_total"))
    total = total if total is not None else parcial
    costo, plazo = num(r.get("costo")), num(r.get("plazo"))
    ejec = gastos = tot = None
    if total is not None and costo is not None:
        if mantenimiento == "Rutinario" and plazo is not None:
            ejec = redondear(total * costo * plazo / 30)
        elif mantenimiento == "Periódico":
            ejec = redondear(total * costo)
    if ejec is not None:
        gastos = redondear(Decimal(ejec) * GASTOS_PCT)
        tot = ejec + gastos
    return {"por_prog": por_prog, "parcial": parcial, "total": total, "ejec": ejec, "gastos": gastos, "tot": tot}


def limpiar(df):
    """Normaliza lo escrito en la tabla y descarta filas totalmente en blanco."""
    d = df.copy()
    for c in COLS_TRAMO:
        if c not in d.columns:
            d[c] = None
        if c in NUM_COLS:
            d[c] = pd.to_numeric(d[c], errors="coerce")
        else:
            d[c] = d[c].apply(lambda v: "" if vacio(v) else str(v).strip())
    for c in ("codigo_ruta", "distrito", "denom_sinac", "denom_interv", "provincia", "ubigeo"):
        d[c] = d[c].str.upper()
    for c in ("prog_ini", "prog_fin"):
        d[c] = d[c].apply(lambda v: fmt_prog(parse_prog(v)) if parse_prog(v) is not None else v)
    util = [c for c in COLS_TRAMO if c != "zona"]
    d = d[~d[util].apply(lambda r: all(vacio(x) for x in r), axis=1)]
    return d.reset_index(drop=True)[COLS_TRAMO]


def completar_longitudes(df, mantenimiento):
    """Antes de guardar: si faltan, la longitud parcial sale de las progresivas y la total = parcial."""
    d = df.copy()
    for i, r in d.iterrows():
        c = calcular_tramo(r, mantenimiento)
        if vacio(r["long_parcial"]) and c["parcial"] is not None:
            d.at[i, "long_parcial"] = float(c["parcial"])
        if vacio(r["long_total"]) and c["total"] is not None:
            d.at[i, "long_total"] = float(c["total"])
    return d


def validar(df, lote):
    """Devuelve (errores, avisos): listas de (n.° de tramo, código de ruta, mensaje).
       Los errores impiden el envío; los avisos se muestran a Sede Central para su revisión."""
    errores, avisos = [], []
    mant, tipo_g = lote["mantenimiento"], lote["tipo_gobierno"]
    cod_depto = UBIGEO_DEPTO.get(str(lote["departamento"]).strip().upper())
    vistos = {}
    for i, r in df.reset_index(drop=True).iterrows():
        n, ruta = i + 1, (r["codigo_ruta"] or "s/código")
        e = lambda m: errores.append((n, ruta, m))  # noqa: E731
        a = lambda m: avisos.append((n, ruta, m))  # noqa: E731
        obligatorios = [("codigo_ruta", "código de ruta"), ("distrito", "distrito"), ("ubigeo", "ubigeo"),
                        ("denom_interv", "denominación de intervención"), ("region_natural", "región natural"),
                        ("estado_cons", "estado de conservación"), ("superficie", "tipo de superficie")]
        if tipo_g == "Regional":
            obligatorios.append(("provincia", "provincia"))
        for k, etiqueta in obligatorios:
            if vacio(r[k]):
                e(f"Falta {etiqueta}.")
        # ubigeo
        if not vacio(r["ubigeo"]):
            codigos = [x.strip() for x in re.split(r"[/,;]", str(r["ubigeo"])) if x.strip()]
            if not codigos or any(not re.fullmatch(r"\d{6}", x) for x in codigos):
                e("El ubigeo debe tener 6 dígitos (si cruza dos distritos: 180211 / 180210).")
            elif cod_depto and any(not x.startswith(cod_depto) for x in codigos):
                a(f"El ubigeo no corresponde al departamento {lote['departamento']} (debería empezar con {cod_depto}).")
        # progresivas y longitudes
        ini, fin = parse_prog(r["prog_ini"]), parse_prog(r["prog_fin"])
        if ini is None or fin is None:
            e("Falta o es inválida la progresiva inicial/final (formato KM XX+XXX, ej. 05+498).")
        elif fin <= ini:
            e("La progresiva final debe ser mayor que la inicial.")
        c = calcular_tramo(r, mant)
        if c["total"] is None or c["total"] <= 0:
            e("Falta la longitud (o las progresivas para calcularla).")
        if c["por_prog"] is not None and c["parcial"] is not None and abs(c["parcial"] - c["por_prog"]) > Decimal("0.0005"):
            a(f"Las progresivas dan {c['por_prog']:.3f} km pero la longitud parcial es {c['parcial']:.3f} km. "
              "Verifique o justifique en Observaciones.")
        if c["total"] is not None and c["parcial"] is not None and c["total"] < c["parcial"]:
            a("La longitud total es menor que la longitud parcial.")
        # costo y plazo
        if vacio(r["costo"]) or (num(r["costo"]) or 0) <= 0:
            e("Falta el costo (Km/Mes en rutinario; Km en periódico).")
        if vacio(r["plazo"]) or (num(r["plazo"]) or 0) <= 0:
            e("Falta el plazo (días calendario).")
        # recomendados
        if vacio(r["denom_sinac"]):
            a("Falta la denominación según SINAC - PVPP.")
        if vacio(r["ancho"]):
            a("Falta el ancho de calzada.")
        for k_x, k_y, lado in (("x_ini", "y_ini", "inicio"), ("x_fin", "y_fin", "final")):
            x, y = num(r[k_x]), num(r[k_y])
            if x is None or y is None:
                a(f"Faltan las coordenadas de {lado}.")
            elif not (Decimal(100000) <= x <= Decimal(900000) and Decimal(7800000) <= y <= Decimal(10200000)):
                a(f"Las coordenadas de {lado} parecen fuera del rango UTM del Perú (X={x}, Y={y}).")
        # duplicados / traslapes
        if ini is not None and fin is not None and not vacio(r["codigo_ruta"]):
            for (n0, ini0, fin0) in vistos.get(r["codigo_ruta"], []):
                if (ini, fin) == (ini0, fin0):
                    e(f"Tramo duplicado con el N° {n0}.")
                elif ini < fin0 and ini0 < fin:
                    a(f"Las progresivas se traslapan con las del tramo N° {n0} de la misma ruta.")
            vistos.setdefault(r["codigo_ruta"], []).append((n, ini, fin))
    return errores, avisos


def tabla_resultados(df, mantenimiento):
    filas = []
    for i, r in df.reset_index(drop=True).iterrows():
        c = calcular_tramo(r, mantenimiento)
        filas.append({"N°": i + 1, "Código de ruta": r["codigo_ruta"],
                      "Grupo": "Pavimentado" if es_pavimentado(r["superficie"]) else "No pavimentado",
                      "Long. por progresivas (km)": None if c["por_prog"] is None else float(c["por_prog"]),
                      "Long. total (km)": None if c["total"] is None else float(c["total"]),
                      "Ejecución (S/)": c["ejec"], "Gastos (S/)": c["gastos"], "Total (S/)": c["tot"]})
    return pd.DataFrame(filas)


def df_vacio():
    fila = {c: (19.0 if c == "zona" else (float("nan") if c in NUM_COLS else "")) for c in COLS_TRAMO}
    return pd.DataFrame([fila])


def _tipificar(df):
    d = df.copy()
    for c in COLS_TRAMO:
        if c in NUM_COLS:
            d[c] = pd.to_numeric(d[c], errors="coerce")
        else:
            d[c] = d[c].astype(str).replace({"nan": "", "None": ""})
    return d[COLS_TRAMO].reset_index(drop=True)


# ---------------------------------------------------------------- Excel (plantilla original)
_SPECS = {
    ("Rutinario", "det"): {
        "hoja": "MR PIA 2025 -ANEXO 1A", "c0": 2, "c1": 30, "label": (5, 11),
        "map": {"B": "nro", "C": "pliego", "D": "region_natural", "E": "departamento", "F": "provincia",
                "G": "distrito", "H": "ubigeo", "I": "codigo_ruta", "J": "denom_sinac", "K": "denom_interv",
                "L": "zona", "M": "prog_ini", "N": "x_ini", "O": "y_ini", "P": "prog_fin", "Q": "x_fin",
                "R": "y_fin", "S": "long_parcial", "T": "long_total", "U": "estado_cons", "V": "superficie",
                "W": "ancho", "X": "tipo_mant", "Y": "costo", "Z": "plazo", "AD": "observaciones"},
        "formulas": {"AA": "=ROUND(T{r}*Y{r}*Z{r}/30,0)", "AB": "=ROUND(AA{r}*0.1,0)", "AC": "=AA{r}+AB{r}"},
        "sumas": ["S", "T", "AA", "AB", "AC"]},
    ("Periódico", "det"): {
        "hoja": "MP PIA 2025 -ANEXO 1B", "c0": 2, "c1": 30, "label": (5, 11),
        "map": {"B": "nro", "C": "pliego", "D": "region_natural", "E": "departamento", "F": "provincia",
                "G": "distrito", "H": "ubigeo", "I": "codigo_ruta", "J": "denom_sinac", "K": "denom_interv",
                "L": "zona", "M": "prog_ini", "N": "x_ini", "O": "y_ini", "P": "prog_fin", "Q": "x_fin",
                "R": "y_fin", "S": "long_parcial", "T": "long_total", "U": "estado_cons", "V": "superficie",
                "W": "ancho", "X": "tipo_mant", "Y": "costo", "Z": "plazo", "AD": "observaciones"},
        "formulas": {"AA": "=ROUND(T{r}*Y{r},0)", "AB": "=ROUND(AA{r}*0.1,0)", "AC": "=AA{r}+AB{r}"},
        "sumas": ["S", "T", "AA", "AB", "AC"]},
    ("Rutinario", "res"): {
        "hoja": "MR PIA 2025 ANEXO 1A", "c0": 2, "c1": 15, "label": (4, 8),
        "map": {"B": "nro", "C": "region_natural", "D": "departamento", "E": "provincia", "F": "distrito",
                "G": "codigo_ruta", "H": "denom_interv", "I": "long_total", "J": "costo", "K": "plazo",
                "O": "observaciones"},
        "formulas": {"L": "=ROUND(I{r}*J{r}*K{r}/30,0)", "M": "=ROUND(L{r}*0.1,0)", "N": "=L{r}+M{r}"},
        "sumas": ["I", "L", "M", "N"]},
    ("Periódico", "res"): {
        "hoja": "MP PIA 2025 ANEXO 1B", "c0": 2, "c1": 14, "label": (4, 8),
        "map": {"B": "nro", "C": "region_natural", "D": "departamento", "E": "provincia", "F": "distrito",
                "G": "codigo_ruta", "H": "denom_interv", "I": "long_total", "J": "costo", "N": "observaciones"},
        "formulas": {"K": "=ROUND(I{r}*J{r},0)", "L": "=ROUND(K{r}*0.1,0)", "M": "=K{r}+L{r}"},
        "sumas": ["I", "K", "L", "M"]},
}


def _col_idx(letra):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letra)


def _ancho_col(ws, letra):
    """Ancho de la columna (openpyxl agrupa columnas contiguas con el mismo ancho)."""
    idx = _col_idx(letra)
    for dim in ws.column_dimensions.values():
        if dim.min and dim.max and dim.min <= idx <= dim.max and dim.width:
            return dim.width
    return 8.43


def _llenar_hoja(ws, spec, lote, tramos, anio_txt, ley):
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.utils import get_column_letter

    c0, c1 = spec["c0"], spec["c1"]
    fila_np, fila_p = 8, 14  # filas modelo de la plantilla
    tipos = {"np": fila_np, "tnp": 13, "p": fila_p, "tp": 18, "tprov": 19, "f20": 20, "f21": 21, "f22": 22, "f23": 23}
    snaps = {k: ({c: copy(ws.cell(r, c)._style) for c in range(c0, c1 + 1)}, ws.row_dimensions[r].height)
             for k, r in tipos.items()}

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 8:
            ws.unmerge_cells(str(rng))
    ws.delete_rows(8, max(ws.max_row, 30))
    ws.conditional_formatting = ConditionalFormattingList()

    def estilo(fila, clave):
        estilos, alto = snaps[clave]
        for c in range(c0, c1 + 1):
            ws.cell(fila, c)._style = copy(estilos[c])
        ws.row_dimensions[fila].height = alto

    es_local = lote["tipo_gobierno"] == "Local"
    via = "VECINALES" if es_local else "DEPARTAMENTALES"
    tipo_mant = "RUTINARIO" if lote["mantenimiento"] == "Rutinario" else "PERIODICO"
    grupos = [[t for t in tramos if not es_pavimentado(t["superficie"])],
              [t for t in tramos if es_pavimentado(t["superficie"])]]
    nro = 0
    fila = 8
    inicio_g, fin_g, tot_g = [], [], []
    for gi, grupo in enumerate(grupos):
        inicio_g.append(fila)
        for t in grupo or [None]:  # al menos una fila (en blanco) por grupo, como en la plantilla
            estilo(fila, "np")
            if t is not None:
                nro += 1
                valores = dict(t, nro=nro, pliego=lote["pliego"], departamento=str(lote["departamento"]).upper(),
                               tipo_mant=tipo_mant)
                if es_local:
                    valores["provincia"] = str(lote["provincia"]).upper()
                for letra, clave in spec["map"].items():
                    v = valores.get(clave)
                    if vacio(v):
                        continue
                    ws[f"{letra}{fila}"] = float(v) if clave in NUM_COLS else v
                for letra, f in spec["formulas"].items():
                    ws[f"{letra}{fila}"] = f.format(r=fila)
                lineas, tam_max = 1, 10
                for letra in spec["map"]:
                    cel = ws[f"{letra}{fila}"]
                    if isinstance(cel.value, str):
                        tam = cel.font.sz or 11
                        tam_max = max(tam_max, tam)
                        cap = max(_ancho_col(ws, letra) * 11 / tam * 0.85, 4)  # caracteres por línea
                        lineas = max(lineas, -(-len(cel.value) // int(cap)))
                ws.row_dimensions[fila].height = min(max(30, lineas * tam_max * 1.35 + 4), 200)
            fila += 1
        fin_g.append(fila - 1)
        estilo(fila, "tnp" if gi == 0 else "tp")
        etiqueta = f"TOTAL CAMINOS {via} {'NO ' if gi == 0 else ''}PAVIMENTADOS"
        ws.cell(fila, spec["label"][0], etiqueta)
        ws.merge_cells(start_row=fila, start_column=spec["label"][0], end_row=fila, end_column=spec["label"][1])
        for letra in spec["sumas"]:
            ws[f"{letra}{fila}"] = f"=SUM({letra}{inicio_g[gi]}:{letra}{fin_g[gi]})"
        tot_g.append(fila)
        fila += 1
    estilo(fila, "tprov")
    ws.cell(fila, spec["label"][0],
            f"TOTAL PROVINCIA DE {str(lote['provincia']).upper()}" if es_local
            else f"TOTAL GOBIERNO REGIONAL DE {str(lote['departamento']).upper()}")
    ws.merge_cells(start_row=fila, start_column=spec["label"][0], end_row=fila, end_column=spec["label"][1])
    for letra in spec["sumas"]:
        ws[f"{letra}{fila}"] = f"=+{letra}{tot_g[0]}+{letra}{tot_g[1]}"
    fila_total = fila
    for k, off in (("f20", 1), ("f21", 2), ("f22", 3), ("f23", 4)):
        estilo(fila_total + off, k)

    # Textos con marcadores (XXXX / 20XX) de la plantilla
    anio = str(lote["anio"])
    for fila_t in range(1, 8):
        for c in range(c0, c1 + 1):
            cel = ws.cell(fila_t, c)
            if isinstance(cel.value, str):
                s = cel.value
                if anio_txt:
                    s = re.sub(r"“Año[^”]*”", f"“{anio_txt}”", s)
                s = s.replace("20XX", anio)
                s = s.replace("VECINALES/DEPARTAMENTALES", via)
                if ley:
                    s = re.sub(r"LEY N°\s*X+", f"LEY N° {ley}", s)
                cel.value = s
    # Área de impresión: mismas columnas, hasta la fila de totales
    m = re.search(r"\$([A-Z]+)\$1:\$([A-Z]+)\$\d+", ws.print_area or "")
    if m:
        ws.print_area = f"{m.group(1)}1:{m.group(2)}{fila_total}"


def generar_anexo(lote, tramos, anio_txt="", ley="", plantilla=PLANTILLA):
    """Devuelve los bytes del Excel (Anexo I detallado + resumen) con el formato de la plantilla original."""
    import openpyxl

    wb = openpyxl.load_workbook(plantilla)
    mant = lote["mantenimiento"]
    conservar = [_SPECS[(mant, "det")]["hoja"], _SPECS[(mant, "res")]["hoja"]]
    for ws in list(wb.worksheets):
        if ws.title not in conservar:
            wb.remove(ws)
    wb._external_links = []  # evita el aviso de "vínculos externos" al abrir el archivo
    for nombre in list(wb.defined_names.keys()):
        del wb.defined_names[nombre]
    for ws in wb.worksheets:
        ws._images = [im for im in ws._images if getattr(im, "width", 0) > 2]  # quita imágenes de 1x1 px
    for tipo in ("det", "res"):
        spec = _SPECS[(mant, tipo)]
        _llenar_hoja(wb[spec["hoja"]], spec, lote, tramos, anio_txt, ley)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@st.cache_data(show_spinner=False, max_entries=20)
def _excel_cache(lote_json, tramos_json, anio_txt, ley):
    return generar_anexo(json.loads(lote_json), json.loads(tramos_json), anio_txt, ley)


# ---------------------------------------------------------------- Google Sheets
@st.cache_resource(show_spinner=False)
def _ws(_libro, nombre):
    cols = LOTES_COLS if nombre == HOJA_LOTES else TRAMOS_COLS
    try:
        return _libro.worksheet(nombre)
    except gspread.WorksheetNotFound:
        pass
    try:
        ws = _libro.add_worksheet(title=nombre, rows=500, cols=len(cols))
        ws.update(range_name="A1", values=[cols], value_input_option="RAW")
        return ws
    except Exception:
        return _libro.worksheet(nombre)


@st.cache_data(ttl=20, show_spinner=False)
def leer(_libro, nombre):
    cols = LOTES_COLS if nombre == HOJA_LOTES else TRAMOS_COLS
    vals = _ws(_libro, nombre).get_all_values()
    if len(vals) < 2:
        df = pd.DataFrame(columns=cols)
        df["_fila"] = []
        return df
    enc = vals[0]
    df = pd.DataFrame([r + [""] * (len(enc) - len(r)) for r in vals[1:]], columns=enc)
    df["_fila"] = range(2, len(df) + 2)
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    return df


def _celda(v):
    if vacio(v):
        return ""
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        f = float(v)
        return int(f) if f.is_integer() else f
    return str(v).strip()


def crear_lote(libro, datos):
    ws = _ws(libro, HOJA_LOTES)
    lote_id = "LT-" + ahora().strftime("%y%m%d-%H%M%S") + "-%02X" % random.randint(0, 255)
    fila = dict(datos, lote_id=lote_id, estado=ESTADO_ELAB, fecha_creacion=ts(), actualizado=ts())
    ws.append_row([_celda(fila.get(c, "")) for c in LOTES_COLS], value_input_option="RAW")
    leer.clear()
    return lote_id


def actualizar_lote(libro, lote_id, cambios):
    ws = _ws(libro, HOJA_LOTES)
    fila = next((i for i, v in enumerate(ws.col_values(1), 1) if i > 1 and v == lote_id), None)
    if not fila:
        raise ValueError("No se encontró la lista de tramos.")
    cambios = dict(cambios, actualizado=ts())
    ws.batch_update([{"range": gspread.utils.rowcol_to_a1(fila, LOTES_COLS.index(k) + 1), "values": [[_celda(v)]]}
                     for k, v in cambios.items()], value_input_option="RAW")
    leer.clear()


def guardar_tramos(libro, lote_id, df):
    """Reemplaza los tramos de la lista. No borra filas de la hoja (marca activo=0), para no desordenar
    a otros usuarios que escriban al mismo tiempo."""
    ws = _ws(libro, HOJA_TRAMOS)
    vals = ws.get_all_values()
    existentes = [i for i, r in enumerate(vals[1:], start=2)
                  if len(r) > 1 and r[0] == lote_id and r[1] == "1"]
    nuevas = [[lote_id, "1", n] + [_celda(r[c]) for c in COLS_TRAMO] for n, (_, r) in enumerate(df.iterrows(), 1)]
    ult = re.sub(r"\d", "", gspread.utils.rowcol_to_a1(1, len(TRAMOS_COLS)))
    cambios = []
    for k, fila in enumerate(existentes):
        if k < len(nuevas):
            cambios.append({"range": f"A{fila}:{ult}{fila}", "values": [nuevas[k]]})
        else:
            cambios.append({"range": f"B{fila}", "values": [["0"]]})
    if cambios:
        ws.batch_update(cambios, value_input_option="RAW")
    extra = nuevas[len(existentes):]
    if extra:
        ws.append_rows(extra, value_input_option="RAW")
    leer.clear()


def tramos_de(libro, lote_id):
    t = leer(libro, HOJA_TRAMOS)
    t = t[(t["lote_id"] == lote_id) & (t["activo"] == "1")].copy()
    t["_n"] = pd.to_numeric(t["nro"], errors="coerce")
    return _tipificar(t.sort_values("_n")[COLS_TRAMO])


def _visibles(lotes, user, usr_key):
    lotes = lotes[lotes["estado"] != ESTADO_ANUL]
    if user["rol"] == "SEDE_CENTRAL" or lotes.empty:
        return lotes
    provs = {p.strip().upper() for p in user["provincias"]}
    regs = {r.strip().upper() for r in user["regiones"]}
    ok = (((lotes["tipo_gobierno"] == "Local") & lotes["provincia"].str.strip().str.upper().isin(provs))
          | ((lotes["tipo_gobierno"] == "Regional") & lotes["departamento"].str.strip().str.upper().isin(regs))
          | (lotes["creado_por"] == usr_key))
    return lotes[ok]


def alertas_usuario(libro, user, usr_key):
    """Avisos para mostrar en la parte superior de la app: lista de (nivel, texto)."""
    lotes = _visibles(leer(libro, HOJA_LOTES), user, usr_key)
    out = []
    if lotes.empty:
        return out
    if user["rol"] == "SEDE_CENTRAL":
        n = int((lotes["estado"] == ESTADO_ENVIADO).sum())
        if n:
            out.append(("warning", f"🔔 Hay {n} lista(s) de tramos enviadas por las unidades zonales que esperan su "
                                   "revisión (pestaña Módulo 2)."))
        return out
    limite = (ahora() - timedelta(days=7)).strftime("%Y-%m-%d %H:%M:%S")
    for _, l in lotes.iterrows():
        etiqueta = f"{l['pliego']} ({l['mantenimiento']} {l['anio']})"
        if l["estado"] == ESTADO_OBS:
            out.append(("error", f"↩️ Sede Central observó la lista de tramos «{etiqueta}». Revise y vuelva a enviarla "
                                 "(Módulo 2)."))
        elif l["estado"] == ESTADO_APROB and l["fecha_revision"] >= limite:
            out.append(("success", f"✅ La lista de tramos «{etiqueta}» fue aprobada: ya puede descargar el Anexo I "
                                   "(Módulo 2)."))
    return out


# ---------------------------------------------------------------- interfaz
def _banner(titulo):
    st.markdown(f'<div class="header-container"><div class="header-title">{titulo}</div></div>',
                unsafe_allow_html=True)


def _flash(msg):
    st.session_state["flash"] = msg
    leer.clear()
    st.rerun()


def _fmt(v, dec=0):
    return "" if v is None else (f"{v:,.{dec}f}")


def _mostrar_control(errores, avisos):
    if errores:
        st.error(f"{len(errores)} error(es) que deben corregirse antes de enviar:")
        for n, ruta, m in errores:
            st.markdown(f"- ❌ **Tramo {n} ({ruta})**: {m}")
    if avisos:
        with st.expander(f"⚠️ {len(avisos)} aviso(s) para revisar", expanded=not errores):
            for n, ruta, m in avisos:
                st.markdown(f"- ⚠️ **Tramo {n} ({ruta})**: {m}")
    if not errores and not avisos:
        st.success("Sin observaciones de control.")


def _mostrar_resultados(d, mant):
    res = tabla_resultados(d, mant)
    if res.empty:
        return
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Longitud total (km)", _fmt(res["Long. total (km)"].sum(), 3))
    c2.metric("Ejecución (S/)", _fmt(res["Ejecución (S/)"].sum()))
    c3.metric("Gastos operativos / supervisión (S/)", _fmt(res["Gastos (S/)"].sum()))
    c4.metric("Total (S/)", _fmt(res["Total (S/)"].sum()))
    vista = res.copy()
    for c in ("Ejecución (S/)", "Gastos (S/)", "Total (S/)"):
        vista[c] = vista[c].apply(lambda v: "" if pd.isna(v) else f"{int(v):,}")
    for c in ("Long. por progresivas (km)", "Long. total (km)"):
        vista[c] = vista[c].apply(lambda v: "" if pd.isna(v) else f"{v:,.3f}")
    st.dataframe(vista, hide_index=True, use_container_width=True)


def _config_editor(es_regional):
    cc = st.column_config
    cfg = {
        "codigo_ruta": cc.TextColumn("Código de ruta", required=True),
        "distrito": cc.TextColumn("Distrito", required=True),
        "ubigeo": cc.TextColumn("Ubigeo", help="6 dígitos. Si cruza dos distritos: 180211 / 180210"),
        "denom_sinac": cc.TextColumn("Denominación según SINAC - PVPP"),
        "denom_interv": cc.TextColumn("Denominación de intervención", required=True),
        "region_natural": cc.SelectboxColumn("Región natural", options=REGIONES_NATURALES),
        "provincia": cc.TextColumn("Provincia"),
        "zona": cc.NumberColumn("Zona", step=1, format="%d"),
        "prog_ini": cc.TextColumn("Progresiva inicial (KM XX+XXX)", help="Ej. 05+498. También acepta 5498 (metros)."),
        "x_ini": cc.NumberColumn("Coord. X inicio", format="%.3f"),
        "y_ini": cc.NumberColumn("Coord. Y inicio", format="%.3f"),
        "prog_fin": cc.TextColumn("Progresiva final (KM XX+XXX)"),
        "x_fin": cc.NumberColumn("Coord. X final", format="%.3f"),
        "y_fin": cc.NumberColumn("Coord. Y final", format="%.3f"),
        "long_parcial": cc.NumberColumn("Longitud parcial (km)", format="%.3f",
                                        help="Si la deja vacía se calcula con las progresivas."),
        "long_total": cc.NumberColumn("Longitud total (km)", format="%.3f", help="Si la deja vacía será igual a la parcial."),
        "estado_cons": cc.SelectboxColumn("Estado de conservación", options=ESTADOS_CONSERVACION),
        "superficie": cc.SelectboxColumn("Tipo de superficie de rodadura", options=SUPERFICIES_NO_PAV + SUPERFICIES_PAV),
        "ancho": cc.NumberColumn("Ancho de calzada (m)", format="%.2f"),
        "costo": cc.NumberColumn("Costo", format="%.2f", help="Rutinario: Km/Mes. Periódico: Km."),
        "plazo": cc.NumberColumn("Plazo (días calendario)", step=1, format="%d"),
        "observaciones": cc.TextColumn("Observaciones"),
    }
    orden = [c for c in COLS_TRAMO if es_regional or c != "provincia"]
    return cfg, orden


def _panel_info(lote):
    c1, c2, c3 = st.columns(3)
    c1.markdown(f"**Gobierno:** {lote['pliego']}  \n**Tipo:** {lote['tipo_gobierno']} · {lote['mantenimiento']}")
    ambito = lote["departamento"] + (f" / {lote['provincia']}" if lote["provincia"] else "")
    c2.markdown(f"**Año:** {lote['anio']}  \n**Ámbito:** {ambito}")
    c3.markdown(f"**Estado:** {lote['estado']}  \n**HT vinculada:** {lote['ht'] or '—'}")
    if lote["estado"] == ESTADO_OBS and lote["obs_central"]:
        st.warning(f"Observación de Sede Central: {lote['obs_central']}")


def _panel_edicion(libro, lote, usr_key):
    lote_id, mant = lote["lote_id"], lote["mantenimiento"]
    clave_base = f"m2_base_{lote_id}_{lote['actualizado']}"
    if clave_base not in st.session_state:
        t = tramos_de(libro, lote_id)
        st.session_state[clave_base] = t if len(t) else df_vacio()
    cfg, orden = _config_editor(lote["tipo_gobierno"] == "Regional")
    st.caption("Complete la tabla (también puede pegar filas copiadas desde Excel). Agregue tramos con la fila vacía "
               "del final. Guarde antes de cambiar de lista.")
    editado = st.data_editor(st.session_state[clave_base], num_rows="dynamic", column_config=cfg,
                             column_order=orden, hide_index=True, use_container_width=True,
                             key=f"m2_ed_{lote_id}_{lote['actualizado']}")
    d = limpiar(editado)
    errores, avisos = validar(d, lote) if len(d) else ([], [])
    st.markdown("##### Resultado de los cálculos")
    _mostrar_resultados(d, mant)
    st.markdown("##### Control de consistencia")
    if len(d):
        _mostrar_control(errores, avisos)
    else:
        st.info("Aún no hay tramos registrados.")

    c1, c2 = st.columns(2)
    if c1.button("💾 Guardar borrador", key=f"m2_guardar_{lote_id}", use_container_width=True):
        try:
            guardar_tramos(libro, lote_id, completar_longitudes(d, mant))
            actualizar_lote(libro, lote_id, {})
        except Exception as e:
            st.error(f"No se pudo guardar: {e}")
        else:
            _flash("Borrador guardado.")
    enviar_ok = len(d) > 0 and not errores
    if c2.button("📤 Enviar a Sede Central", key=f"m2_enviar_{lote_id}", type="primary", use_container_width=True,
                 disabled=not enviar_ok, help=None if enviar_ok else "Corrija los errores y registre al menos un tramo."):
        try:
            guardar_tramos(libro, lote_id, completar_longitudes(d, mant))
            actualizar_lote(libro, lote_id, {"estado": ESTADO_ENVIADO, "fecha_envio": ts(), "obs_central": ""})
        except Exception as e:
            st.error(f"No se pudo enviar: {e}")
        else:
            _flash("Lista enviada a Sede Central para su revisión.")
    if lote["estado"] == ESTADO_ELAB:
        with st.expander("🗑️ Anular esta lista"):
            if st.checkbox("Confirmo que deseo anular la lista", key=f"m2_anul_c_{lote_id}") and \
                    st.button("Anular lista", key=f"m2_anul_{lote_id}"):
                actualizar_lote(libro, lote_id, {"estado": ESTADO_ANUL})
                _flash("La lista fue anulada.")


def _panel_lectura(libro, lote, con_revision, usr_key):
    lote_id, mant = lote["lote_id"], lote["mantenimiento"]
    t = tramos_de(libro, lote_id)
    d = limpiar(t) if len(t) else t
    st.dataframe(d.rename(columns={c: c.replace("_", " ").capitalize() for c in COLS_TRAMO}), hide_index=True,
                 use_container_width=True)
    st.markdown("##### Resultado de los cálculos")
    _mostrar_resultados(d, mant)
    errores, avisos = validar(d, lote) if len(d) else ([], [])
    st.markdown("##### Control de consistencia")
    _mostrar_control(errores, avisos)
    if con_revision and lote["estado"] == ESTADO_ENVIADO:
        st.markdown("##### Revisión de Sede Central")
        obs = st.text_area("Observación (obligatoria si va a observar la lista)", key=f"m2_obs_{lote_id}")
        c1, c2 = st.columns(2)
        if c1.button("✅ Aprobar lista", key=f"m2_aprobar_{lote_id}", type="primary", use_container_width=True,
                     disabled=bool(errores)):
            actualizar_lote(libro, lote_id, {"estado": ESTADO_APROB, "revisado_por": usr_key,
                                             "fecha_revision": ts(), "obs_central": obs.strip()})
            _flash("Lista aprobada. El especialista zonal ya puede descargar el Anexo I.")
        if c2.button("↩️ Observar y devolver", key=f"m2_observar_{lote_id}", use_container_width=True):
            if not obs.strip():
                st.error("Escriba la observación para devolver la lista al especialista zonal.")
            else:
                actualizar_lote(libro, lote_id, {"estado": ESTADO_OBS, "revisado_por": usr_key,
                                                 "fecha_revision": ts(), "obs_central": obs.strip()})
                _flash("Lista observada y devuelta al especialista zonal.")
    if lote["estado"] == ESTADO_APROB:
        _panel_descarga(lote, d)


def _panel_descarga(lote, d):
    letra = "I-A" if lote["mantenimiento"] == "Rutinario" else "I-B"
    st.markdown(f"##### Descargar Anexo {letra} y su resumen")
    c1, c2 = st.columns(2)
    anio_txt = c1.text_input("Denominación oficial del año (opcional)",
                             value=ANIO_DEFECTO.get(int(lote["anio"] or 0), ""), key=f"m2_anio_{lote['lote_id']}")
    ley = c2.text_input("Ley N° (opcional)", value=lote["ley"], key=f"m2_ley_d_{lote['lote_id']}")
    if not PLANTILLA.exists():
        st.error("No se encontró la plantilla 'plantilla_anexo_i.xlsx' junto a app.py.")
        return
    tramos = json.loads(d.to_json(orient="records", force_ascii=False))
    datos = _excel_cache(json.dumps(lote, ensure_ascii=False, default=str),
                         json.dumps(tramos, ensure_ascii=False), anio_txt.strip(), ley.strip())
    nombre = re.sub(r"[^A-Za-z0-9]+", "_", f"ANEXO_{letra}_{lote['pliego']}_{lote['anio']}").strip("_") + ".xlsx"
    st.download_button(f"⬇️ Descargar Anexo {letra} (Excel)", data=datos, file_name=nombre,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")


def render(libro, user, usr_key, hts_visibles):
    es_central = user["rol"] == "SEDE_CENTRAL"
    _banner(TITULO)
    try:
        lotes = leer(libro, HOJA_LOTES)
        leer(libro, HOJA_TRAMOS)
    except Exception as e:
        st.error(f"No se pudo leer la información del Módulo 2: {e}")
        return
    visibles = _visibles(lotes, user, usr_key)

    if not es_central:
        with st.expander("➕ Nueva lista de tramos remitida por un gobierno", expanded=visibles.empty):
            tipo_g = st.radio("Tipo de gobierno", TIPOS_GOBIERNO, horizontal=True, key="m2_tg")
            pliego = st.text_input("Pliego (ej. MUNICIPALIDAD PROVINCIAL DE TACNA)", key="m2_pliego")
            mant = st.radio("Tipo de mantenimiento", ["Rutinario"] if tipo_g == "Local" else MANTENIMIENTOS,
                            horizontal=True, key="m2_mant")
            c1, c2, c3 = st.columns(3)
            anio = c1.number_input("Año del presupuesto", 2024, 2040, ahora().year, key="m2_anio_n")
            depto = c2.selectbox("Departamento", user["regiones"], key="m2_depto")
            prov = c3.selectbox("Provincia", user["provincias"], key="m2_prov") if tipo_g == "Local" else ""
            c4, c5 = st.columns(2)
            ht = c4.selectbox("Solicitud (HT) del Módulo 1 vinculada (opcional)", ["(sin vincular)"] + list(hts_visibles),
                              key="m2_ht")
            ley = c5.text_input("Ley N° (opcional)", key="m2_ley")
            if st.button("Crear lista", key="m2_crear", type="primary"):
                if not pliego.strip():
                    st.error("Ingrese el pliego (nombre del gobierno).")
                else:
                    try:
                        nuevo = crear_lote(libro, {"tipo_gobierno": tipo_g, "pliego": pliego.strip().upper(),
                                                   "mantenimiento": mant, "anio": int(anio), "departamento": depto,
                                                   "provincia": prov, "ht": "" if ht.startswith("(") else ht,
                                                   "ley": ley.strip(), "creado_por": usr_key})
                    except Exception as e:
                        st.error(f"No se pudo crear la lista: {e}")
                    else:
                        st.session_state["m2_sel"] = nuevo
                        _flash("Lista creada. Complete los tramos y envíela a Sede Central.")

    if visibles.empty:
        st.info("Aún no hay listas de tramos registradas.")
        return

    # Resumen por estado
    cuenta = visibles["estado"].value_counts()
    k = st.columns(4)
    for col, est in zip(k, (ESTADO_ELAB, ESTADO_ENVIADO, ESTADO_OBS, ESTADO_APROB)):
        col.metric(est, int(cuenta.get(est, 0)))

    lista = visibles.sort_values("_fila", ascending=False)
    if es_central:
        modo = st.radio("Mostrar", ["Pendientes de revisión", "Todas"], horizontal=True, key="m2_modo")
        if modo == "Pendientes de revisión":
            lista = lista[lista["estado"] == ESTADO_ENVIADO]
    if lista.empty:
        st.success("No hay listas pendientes de revisión.")
        return
    ids = list(lista["lote_id"])
    etiquetas = {r["lote_id"]: f"{r['pliego']} · {r['mantenimiento']} {r['anio']} · {r['estado']}"
                 for _, r in lista.iterrows()}
    if st.session_state.get("m2_sel") not in ids:
        st.session_state.pop("m2_sel", None)
    sel = st.selectbox("Lista de tramos", ids, format_func=lambda i: etiquetas[i], key="m2_sel")
    lote = lista[lista["lote_id"] == sel].iloc[0].to_dict()

    st.markdown("---")
    _panel_info(lote)
    if not es_central and lote["estado"] in (ESTADO_ELAB, ESTADO_OBS):
        _panel_edicion(libro, lote, usr_key)
    else:
        _panel_lectura(libro, lote, es_central, usr_key)
